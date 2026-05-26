#!/usr/bin/env python3
"""
Generate DOCX documentation with all code and screenshots (desktop + mobile).
"""
import os
import sys
import time
import threading
from http.server import HTTPServer, SimpleHTTPRequestHandler
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = r"D:\PROJECTS\WEB-DEV-VOC\ASSIGNMENT2"
OUTPUT = r"D:\PROJECTS\WEB-DEV-VOC\ASSIGNMENT2\ASSIGNMENT2_DOCUMENTATION.docx"
PORT = 8765

SECTIONS = [
    ("Navbar", "#home"),       # navbar is visible on hero
    ("Hero Section", "#home"),
    ("About Me", "#about"),
    ("Skills", "#skills"),
    ("Projects", "#projects"),
    ("Contact Form", "#contact"),
    ("Footer", "#contact"),    # scroll to contact, footer is below
]

# ---- HTTP Server ----
def start_server():
    os.chdir(BASE_DIR)
    server = HTTPServer(("127.0.0.1", PORT), SimpleHTTPRequestHandler)
    server.serve_forever()

# ---- Screenshot helpers ----
def get_driver(width):
    opts = Options()
    opts.add_argument("--headless=new")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-gpu")
    opts.add_argument(f"--window-size={width},1200")
    opts.add_argument("--hide-scrollbars")
    d = webdriver.Chrome(options=opts)
    d.set_window_size(width, 1200)
    return d

def section_screenshot(driver, section_name, anchor, width, ss_dir):
    """Scroll to anchor, wait, then take a full-viewport screenshot."""
    url = f"http://127.0.0.1:{PORT}/index.html"
    driver.get(url)
    time.sleep(1)

    # click anchor via JS to scroll there
    driver.execute_script(f"window.location.hash = '{anchor}';")
    time.sleep(0.8)

    # scroll up a bit to show section top if it's the first
    if anchor == "#home":
        driver.execute_script("window.scrollTo(0, 0);")
        time.sleep(0.3)

    size_label = "desktop" if width >= 1024 else "mobile"
    fname = f"{section_name.replace(' ', '_').lower()}_{size_label}.png"
    fpath = os.path.join(ss_dir, fname)
    driver.save_screenshot(fpath)
    print(f"  -> {fname}")
    return fpath

# ---- DOCX builder ----
def add_code_block(doc, label, code):
    p = doc.add_paragraph()
    run = p.add_run(label)
    run.bold = True
    run.font.size = Pt(12)
    for line in code.strip().split("\n"):
        p = doc.add_paragraph(line, style="No Spacing")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        for run in p.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(8)
    doc.add_paragraph()

def build_docx(ss_dir, desktop_shots, mobile_shots):
    doc = Document()

    # Title section - black & white
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ASSIGNMENT 2 - Personal Portfolio")
    run.bold = True
    run.font.size = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Chaturvedhi Narsimha")
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Web Development Internship")
    run.font.size = Pt(11)

    doc.add_page_break()

    # ---- Code sections ----
    p = doc.add_paragraph()
    run = p.add_run("Source Code")
    run.bold = True
    run.font.size = Pt(16)

    for fname in ["index.html", "style.css", "script.js"]:
        fpath = os.path.join(BASE_DIR, fname)
        with open(fpath, "r", encoding="utf-8") as fh:
            code = fh.read()
        add_code_block(doc, fname, code)

    doc.add_page_break()

    # ---- Screenshots ----
    p = doc.add_paragraph()
    run = p.add_run("Screenshots")
    run.bold = True
    run.font.size = Pt(16)

    for i, (sec_name, _) in enumerate(SECTIONS):
        p = doc.add_paragraph()
        run = p.add_run(sec_name)
        run.bold = True
        run.font.size = Pt(13)

        # desktop
        des = desktop_shots[i]
        p = doc.add_paragraph("Desktop:")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(des, width=Inches(4.5))

        # spacing
        doc.add_paragraph()

        # mobile
        mob = mobile_shots[i]
        p = doc.add_paragraph("Mobile:")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(mob, width=Inches(2.8))

        # spacer between sections
        doc.add_paragraph()

    doc.save(OUTPUT)
    print(f"\nDOCX saved to {OUTPUT}")

# ---- Main ----
def main():
    # create screenshots directory
    ss_dir = os.path.join(BASE_DIR, "_screenshots")
    os.makedirs(ss_dir, exist_ok=True)

    # start server in background
    svr_thread = threading.Thread(target=start_server, daemon=True)
    svr_thread.start()
    time.sleep(1)

    desktop_shots = []
    mobile_shots = []

    print("Taking desktop screenshots (1280px)...")
    d1 = get_driver(1280)
    for sec_name, anchor in SECTIONS:
        p = section_screenshot(d1, sec_name, anchor, 1280, ss_dir)
        desktop_shots.append(p)
    d1.quit()

    print("Taking mobile screenshots (375px)...")
    d2 = get_driver(375)
    for sec_name, anchor in SECTIONS:
        p = section_screenshot(d2, sec_name, anchor, 375, ss_dir)
        mobile_shots.append(p)
    d2.quit()

    print("Building DOCX...")
    build_docx(ss_dir, desktop_shots, mobile_shots)

    print("Done!")
    sys.exit(0)

if __name__ == "__main__":
    main()
